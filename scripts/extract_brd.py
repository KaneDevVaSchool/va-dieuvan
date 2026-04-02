from __future__ import annotations

import json
from pathlib import Path

from docx import Document


def main() -> None:
    # Update this path to extract other BRD versions
    docx_path = Path(r"c:\Users\ASUS\Downloads\brd2.docx")
    if not docx_path.exists():
        raise SystemExit(f"BRD docx not found: {docx_path}")

    out_dir = Path(__file__).resolve().parents[1] / "storage" / "app" / "brd_extract_v2"
    out_dir.mkdir(parents=True, exist_ok=True)

    doc = Document(str(docx_path))

    paragraphs: list[str] = []
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if t:
            paragraphs.append(t)

    tables: list[dict] = []
    for ti, table in enumerate(doc.tables):
        rows: list[list[str]] = []
        for row in table.rows:
            rows.append([((cell.text or "")).strip() for cell in row.cells])
        tables.append({"index": ti, "rows": rows})

    (out_dir / "brd_text.txt").write_text("\n".join(paragraphs), encoding="utf-8")
    (out_dir / "brd_tables.json").write_text(
        json.dumps(tables, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    print(f"Extracted BRD -> {out_dir}")
    print(f"Paragraphs: {len(paragraphs)}")
    print(f"Tables: {len(tables)}")


if __name__ == "__main__":
    main()

